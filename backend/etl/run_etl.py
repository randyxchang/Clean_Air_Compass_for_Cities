#!/usr/bin/env python3
"""
ETL pipeline for Clean Air Compass.
Reads source CSVs/Excel from the data repo and writes processed JSON into backend/data/.

Usage:
    python3 backend/etl/run_etl.py --source /path/to/Clean_Air_Compass_for_Cities-main
"""

import argparse
import csv
import json
import os
import shutil
from pathlib import Path

import openpyxl

SCRIPT_DIR = Path(__file__).parent
BACKEND_DATA = SCRIPT_DIR.parent / "data"
PUBLIC_DIR = SCRIPT_DIR.parent.parent / "public"

# Manual mapping: health impact city names -> uesi_id
HEALTH_CITY_MAP = {
    "Douala": None,  # Not in uesi_city.csv
    "Yaounde": "CMR_R2746229",
    "Addis Ababa": "ETH_R1707699",
    "Nairobi": "KEN_R9185096",
    "Kisumu": None,
    "Accra": None,
    "Kampala": "UGA_R10546821",
    "Jinja": None,
    "Jinja ": None,
    "FortPortal": None,
    "Pretoria": "ZAF_R594508",
    "Free Town": "SLE_N27565056",
    "Bujumbura": None,
    "Luanda": "AGO_R11704226",
    "Lagos": "NGA_R3718182",
    "Port Louis": None,
}


def read_csv(path, encoding="utf-8-sig"):
    with open(path, encoding=encoding, newline="") as f:
        return list(csv.DictReader(f))


def write_json(path, data):
    path.parent.mkdir(parents=True, exist_ok=True)
    with open(path, "w") as f:
        json.dump(data, f, separators=(",", ":"))


def compute_centroid(geojson_path):
    """Compute centroid from GeoJSON bounding box."""
    with open(geojson_path) as f:
        data = json.load(f)

    min_lat, max_lat = 90, -90
    min_lng, max_lng = 180, -180

    def walk_coords(obj):
        nonlocal min_lat, max_lat, min_lng, max_lng
        if isinstance(obj, list):
            if len(obj) >= 2 and isinstance(obj[0], (int, float)) and isinstance(obj[1], (int, float)):
                lng, lat = obj[0], obj[1]
                min_lat = min(min_lat, lat)
                max_lat = max(max_lat, lat)
                min_lng = min(min_lng, lng)
                max_lng = max(max_lng, lng)
            else:
                for item in obj:
                    walk_coords(item)

    for feature in data.get("features", []):
        geom = feature.get("geometry", {})
        walk_coords(geom.get("coordinates", []))

    if min_lat <= max_lat:
        return round((min_lat + max_lat) / 2, 4), round((min_lng + max_lng) / 2, 4)
    return None, None


def simplify_geojson(geojson_path, output_path, precision=5):
    """Copy GeoJSON with coordinate precision reduced."""
    with open(geojson_path) as f:
        data = json.load(f)

    def truncate_coords(obj):
        if isinstance(obj, list):
            if len(obj) >= 2 and isinstance(obj[0], (int, float)) and isinstance(obj[1], (int, float)):
                return [round(obj[0], precision), round(obj[1], precision)] + obj[2:]
            return [truncate_coords(item) for item in obj]
        return obj

    for feature in data.get("features", []):
        geom = feature.get("geometry", {})
        if "coordinates" in geom:
            geom["coordinates"] = truncate_coords(geom["coordinates"])

    output_path.parent.mkdir(parents=True, exist_ok=True)
    with open(output_path, "w") as f:
        json.dump(data, f, separators=(",", ":"))


def etl_cities(source_dir):
    """Step 1a: Build master cities.json."""
    print("  Loading uesi_city.csv...")
    city_lookup = {}
    for row in read_csv(source_dir / "data" / "uesi_city.csv"):
        city_lookup[row["uesi_id"]] = {
            "name": row["city"],
            "iso": row["iso"],
        }

    # Load equity data for region/tier
    print("  Loading equity data...")
    equity_data = {}
    for row in read_csv(source_dir / "data" / "equity" / "eq_city.csv"):
        equity_data[row["uesi_id"]] = {
            "region": row.get("region", ""),
            "tier": row.get("Tier", ""),
            "PM25_INC_GINI": safe_float(row.get("PM25_INC_GINI")),
            "NO2_INC_GINI": safe_float(row.get("NO2_INC_GINI")),
            "PM25_CONC": safe_float(row.get("PM25_CONC")),
            "NO2_CONC": safe_float(row.get("NO2_CONC")),
        }

    # Load composite index — POL cities first (they override NOPOL)
    print("  Loading composite index (POL)...")
    composite = {}
    pol_cities = set()
    for row in read_csv(source_dir / "data" / "composite_index" / "AQ_FINAL_COMP_POL_APR24.csv"):
        uid = row["uesi_id"]
        pol_cities.add(uid)
        composite[uid] = parse_composite_row(row, has_policy=True)

    print("  Loading composite index (NOPOL)...")
    for row in read_csv(source_dir / "data" / "composite_index" / "AQ_FINAL_COMP_NOPOL_APR24.csv"):
        uid = row["uesi_id"]
        if uid not in composite:
            composite[uid] = parse_composite_row(row, has_policy=False)

    # Check which cities have smogstripes
    print("  Checking smogstripes availability...")
    smogstripes_ids = set()
    for row in read_csv(source_dir / "data" / "smogstripes" / "data_shiny.csv"):
        smogstripes_ids.add(row["uesi_id"])

    # Check which cities have GeoJSON
    geojson_dir = source_dir / "data" / "geometry"
    geojson_ids = set()
    if geojson_dir.exists():
        for f in geojson_dir.glob("*.geojson"):
            geojson_ids.add(f.stem)

    # Check which cities have daily data
    print("  Checking daily data availability...")
    daily_ids = set()
    with open(source_dir / "data" / "smogstripes" / "data_shiny_daily.csv", encoding="utf-8-sig") as f:
        reader = csv.DictReader(f)
        for row in reader:
            daily_ids.add(row["uesi_id"])

    # Health impact city IDs
    health_ids = set(v for v in HEALTH_CITY_MAP.values() if v)

    # Check which cities have images
    image_dir = source_dir / "images" / "exports"
    image_ids = set()
    if image_dir.exists():
        for f in image_dir.glob("*_transparent.png"):
            name = f.stem.replace("_transparent", "")
            if " " not in name:
                image_ids.add(name)

    # Compute centroids from GeoJSON
    print("  Computing centroids from GeoJSON...")
    centroids = {}
    for uid in geojson_ids:
        geojson_path = geojson_dir / f"{uid}.geojson"
        lat, lng = compute_centroid(geojson_path)
        if lat is not None:
            centroids[uid] = (lat, lng)

    # Build master list
    print("  Building master cities list...")
    all_uesi_ids = set(city_lookup.keys()) | set(composite.keys())
    cities = []

    for uid in sorted(all_uesi_ids):
        info = city_lookup.get(uid, {"name": uid, "iso": uid[:3]})
        comp = composite.get(uid, {})
        eq = equity_data.get(uid, {})
        lat, lng = centroids.get(uid, (None, None))

        city = {
            "id": uid,
            "name": info["name"],
            "iso": info["iso"],
            "lat": lat,
            "lng": lng,
            "region": eq.get("region", ""),
            "tier": eq.get("tier", ""),
            "pm25_avg": comp.get("pm25_avg"),
            "no2_avg": comp.get("no2_avg"),
            "final_score": comp.get("final_score"),
            "has_policy": uid in pol_cities,
            "has_health": uid in health_ids,
            "has_smogstripes": uid in smogstripes_ids,
            "has_geojson": uid in geojson_ids,
            "has_daily": uid in daily_ids,
            "has_image": uid in image_ids,
        }
        cities.append(city)

    write_json(BACKEND_DATA / "cities.json", cities)
    print(f"  -> cities.json: {len(cities)} cities")
    return cities, composite, equity_data


def parse_composite_row(row, has_policy):
    """Parse a composite index CSV row into a dict."""
    result = {
        "pm25_avg": safe_float(row.get("pm25_avg")),
        "no2_avg": safe_float(row.get("no2_avg")),
        "final_score": safe_float(row.get("final_score")),
        "has_policy": has_policy,
        "scores": {},
        "raw_values": {},
    }

    score_fields = [
        "pm25_conc_score", "pm25_long_trend_score", "pm25_short_trend_score",
        "no2_conc_score", "no2_long_trend_score", "no2_short_trend_score",
        "PM25_eq_score", "NO2_eq_score",
        "pm25_score", "no2_score", "equity_score",
    ]
    for field in score_fields:
        if field in row:
            result["scores"][field] = safe_float(row[field])

    raw_fields = [
        "pm25_avg", "pm25_long_trend_avg", "pm25_short_trend_avg",
        "no2_avg", "no2_long_trend_avg", "no2_short_trend_avg",
        "year_pm25", "year_no2",
    ]
    for field in raw_fields:
        if field in row:
            result["raw_values"][field] = safe_float(row[field])

    if has_policy:
        result["policy_score"] = safe_float(row.get("pol_aq_comp"))

    return result


def etl_city_details(composite, equity_data):
    """Step 1b: Write per-city detail JSON files."""
    print("  Writing per-city detail files...")
    out_dir = BACKEND_DATA / "cities"
    out_dir.mkdir(parents=True, exist_ok=True)

    count = 0
    for uid, comp in composite.items():
        eq = equity_data.get(uid, {})
        detail = {
            "id": uid,
            "scores": comp.get("scores", {}),
            "has_policy": comp.get("has_policy", False),
            "final_score": comp.get("final_score"),
            "raw_values": comp.get("raw_values", {}),
            "equity": eq,
        }
        if comp.get("has_policy"):
            detail["policy_score"] = comp.get("policy_score")

        write_json(out_dir / f"{uid}.json", detail)
        count += 1

    print(f"  -> cities/: {count} files")


def etl_smogstripes(source_dir):
    """Step 1c: Write per-city smogstripes JSON files (2021 WHO only)."""
    print("  Processing smogstripes data...")
    out_dir = BACKEND_DATA / "smogstripes"
    out_dir.mkdir(parents=True, exist_ok=True)

    # Group by city, filter to who_year=2021
    city_data = {}
    for row in read_csv(source_dir / "data" / "smogstripes" / "data_shiny.csv"):
        if row.get("who_year") != "2021":
            continue
        uid = row["uesi_id"]
        if uid not in city_data:
            city_data[uid] = {
                "id": uid,
                "city": row["city"],
                "trend_estimate": safe_float(row.get("estimate")),
                "trend_p_value": safe_float(row.get("p.value")),
                "years": [],
            }
        city_data[uid]["years"].append({
            "year": int(row["year"]),
            "pm25_mean": safe_float(row["pm25_mean"]),
            "who_category": row.get("who_val", ""),
        })

    for uid, data in city_data.items():
        data["years"].sort(key=lambda x: x["year"])
        write_json(out_dir / f"{uid}.json", data)

    print(f"  -> smogstripes/: {len(city_data)} files")


def etl_daily(source_dir):
    """Step 1d: Aggregate daily data to monthly summaries."""
    print("  Processing daily data (this may take a moment)...")
    out_dir = BACKEND_DATA / "daily"
    out_dir.mkdir(parents=True, exist_ok=True)

    # Read and aggregate
    city_monthly = {}  # uid -> month -> list of pm25 values
    city_who_dist = {}
    city_conc_dist = {}
    city_anomaly_dist = {}

    with open(source_dir / "data" / "smogstripes" / "data_shiny_daily.csv", encoding="utf-8-sig") as f:
        reader = csv.DictReader(f)
        for row in reader:
            uid = row["uesi_id"]
            pm25 = safe_float(row.get("pm25"))
            date_str = row.get("date", "")

            if uid not in city_monthly:
                city_monthly[uid] = {}
                city_who_dist[uid] = {}
                city_conc_dist[uid] = {}
                city_anomaly_dist[uid] = {}

            # Monthly aggregation
            if date_str and pm25 is not None:
                try:
                    month = int(date_str.split("-")[1])
                except (IndexError, ValueError):
                    continue
                if month not in city_monthly[uid]:
                    city_monthly[uid][month] = []
                city_monthly[uid][month].append(pm25)

            # Distribution counts
            who_cat = row.get("WHO", "")
            if who_cat:
                city_who_dist[uid][who_cat] = city_who_dist[uid].get(who_cat, 0) + 1

            conc_cat = row.get("Concentration", "")
            if conc_cat:
                city_conc_dist[uid][conc_cat] = city_conc_dist[uid].get(conc_cat, 0) + 1

            anomaly_cat = row.get("Anomaly", "")
            if anomaly_cat:
                city_anomaly_dist[uid][anomaly_cat] = city_anomaly_dist[uid].get(anomaly_cat, 0) + 1

    for uid in city_monthly:
        monthly = []
        for month in sorted(city_monthly[uid].keys()):
            vals = city_monthly[uid][month]
            monthly.append({
                "month": month,
                "avg_pm25": round(sum(vals) / len(vals), 1),
                "max_pm25": round(max(vals), 1),
                "min_pm25": round(min(vals), 1),
                "days": len(vals),
            })

        data = {
            "id": uid,
            "year": 2021,
            "monthly": monthly,
            "who_distribution": city_who_dist.get(uid, {}),
            "concentration_distribution": city_conc_dist.get(uid, {}),
            "anomaly_distribution": city_anomaly_dist.get(uid, {}),
        }
        write_json(out_dir / f"{uid}.json", data)

    print(f"  -> daily/: {len(city_monthly)} files")


def etl_health(source_dir):
    """Step 1e: Process health impact Excel data."""
    print("  Processing health impact data...")
    out_dir = BACKEND_DATA / "health"
    out_dir.mkdir(parents=True, exist_ok=True)

    xlsx_path = source_dir / "data" / "health_impact" / "Countries and indicators.xlsx"
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)

    # Sheet1: source breakdown + health (cols E-Y for sources, Z-AS for health)
    ws1 = wb["Sheet1"]
    # Row 2 has column headers
    headers = [cell.value for cell in ws1[2]]

    source_col_names = [
        "Agriculture Contribution (%)",
        "Energy Coal Contribution (%)",
        "Energy NonCoal Contribution (%)",
        "Industry Coal Contribution (%)",
        "Industry NonCoal Contribution (%)",
        "NonRoad Transport Contribution (%)",
        "Road Transport Contribution (%)",
        "Residential Coal Combustion Contribution (%)",
        "Residential Biofuel Combustion Contribution (%)",
        "Residential Other Combustion Contribution (%)",
        "Commercial Combustion Contribution (%)",
        "Other Combustion Contribution (%)",
        "Solvent Contribution (%)",
        "Waste Contribution (%)",
        "International Shipping Contribution (%)",
        "Agricultural Waste Burning Contribution (%)",
        "Other Open Fire Contribution (%)",
        "AFCID Dust Contribution (%)",
        "Windblown Dust Contribution (%)",
        "Remaining Sources Contribution (%)",
    ]

    # Build source col indices
    source_indices = {}
    for name in source_col_names:
        for i, h in enumerate(headers):
            if h and name.strip() == str(h).strip():
                source_indices[name] = i
                break

    # Find PM2.5 weighted avg column
    pm25_col = None
    for i, h in enumerate(headers):
        if h and "Population Weighted" in str(h):
            pm25_col = i
            break

    # Find total attributable mortality (GBD2019 CRF) — column Z (index 25)
    mortality_col = 25  # GBD2019 CRF total mortality

    sheet1_data = {}
    current_country = None
    for row_idx in range(3, ws1.max_row + 1):
        row_vals = [cell.value for cell in ws1[row_idx]]
        city_name = row_vals[2]
        if not city_name:
            continue

        country = row_vals[1] if row_vals[1] else current_country
        if row_vals[1]:
            current_country = row_vals[1]

        source_breakdown = {}
        for name, idx in source_indices.items():
            val = row_vals[idx]
            source_breakdown[name.replace(" Contribution (%)", "")] = val if val else 0

        sheet1_data[city_name.strip()] = {
            "country": country,
            "pm25_weighted_avg": row_vals[pm25_col] if pm25_col else None,
            "total_mortality_gbd": row_vals[mortality_col] if mortality_col < len(row_vals) else None,
            "source_breakdown": source_breakdown,
        }

    # Sheet2: health outcomes
    ws2 = wb["Sheet2"]
    sheet2_data = {}
    for row_idx in range(2, ws2.max_row + 1):
        row_vals = [cell.value for cell in ws2[row_idx]]
        city_name = row_vals[0]
        if not city_name:
            continue
        sheet2_data[city_name.strip()] = {
            "COPD": row_vals[1],
            "DM": row_vals[2],
            "LRI": row_vals[3],
            "LC": row_vals[4],
            "IHD": row_vals[5],
            "Stroke": row_vals[6],
            "preterm_births": row_vals[7],
            "low_birth_weight": row_vals[8],
        }

    # Life expectancy data
    life_exp = {}
    life_path = source_dir / "data" / "health_impact" / "gain_life_exp_aqli_who_2021.xlsx"
    if life_path.exists():
        wb2 = openpyxl.load_workbook(life_path, data_only=True)
        ws_life = wb2.active
        for row_idx in range(2, ws_life.max_row + 1):
            row_vals = [cell.value for cell in ws_life[row_idx]]
            if row_vals[1]:
                life_exp[row_vals[1].strip()] = row_vals[2]

    # Merge and write per-city files
    count = 0
    all_health_cities = set(list(sheet1_data.keys()) + list(sheet2_data.keys()))
    for city_name in all_health_cities:
        uesi_id = HEALTH_CITY_MAP.get(city_name.strip())
        if not uesi_id:
            continue

        s1 = sheet1_data.get(city_name, {})
        s2 = sheet2_data.get(city_name, {})
        le = life_exp.get(city_name)

        # Build grand categories from source breakdown
        sb = s1.get("source_breakdown", {})
        grand_categories = {
            "Energy & Industry": sum_vals(sb, ["Energy Coal", "Energy NonCoal", "Industry Coal", "Industry NonCoal", "Commercial Combustion"]),
            "Transport": sum_vals(sb, ["Road Transport", "NonRoad Transport", "International Shipping"]),
            "Residential": sum_vals(sb, ["Residential Coal Combustion", "Residential Biofuel Combustion", "Residential Other Combustion"]),
            "Agriculture": sum_vals(sb, ["Agriculture", "Agricultural Waste Burning"]),
            "Waste": sum_vals(sb, ["Waste"]),
            "Other": sum_vals(sb, ["Other Combustion", "Solvent", "Other Open Fire", "AFCID Dust", "Windblown Dust", "Remaining Sources"]),
        }

        result = {
            "id": uesi_id,
            "city": city_name,
            "country": s1.get("country", ""),
            "pm25_weighted_avg": s1.get("pm25_weighted_avg"),
            "total_mortality": s1.get("total_mortality_gbd"),
            "source_breakdown": sb,
            "grand_categories": grand_categories,
            "health_outcomes": s2 if s2 else None,
            "life_expectancy_gain": le,
        }
        write_json(out_dir / f"{uesi_id}.json", result)
        count += 1

    print(f"  -> health/: {count} files")


def etl_geometry(source_dir):
    """Step 1f: Copy and simplify GeoJSON files."""
    print("  Simplifying GeoJSON files...")
    src_dir = source_dir / "data" / "geometry"
    out_dir = BACKEND_DATA / "geometry"
    out_dir.mkdir(parents=True, exist_ok=True)

    count = 0
    for geojson_file in sorted(src_dir.glob("*.geojson")):
        simplify_geojson(geojson_file, out_dir / geojson_file.name)
        count += 1

    print(f"  -> geometry/: {count} files")


def etl_images(source_dir):
    """Step 1g: Copy city map images."""
    print("  Copying city images...")
    src_dir = source_dir / "images" / "exports"
    out_dir = PUBLIC_DIR / "images"
    out_dir.mkdir(parents=True, exist_ok=True)

    count = 0
    for png in src_dir.glob("*_transparent.png"):
        if " " in png.name:
            continue
        shutil.copy2(png, out_dir / png.name)
        count += 1

    print(f"  -> frontend/images/: {count} files")


def sum_vals(d, keys):
    total = 0
    for k in keys:
        v = d.get(k, 0)
        if v and isinstance(v, (int, float)):
            total += v
    return round(total, 1)


def safe_float(val):
    if val is None or val == "":
        return None
    try:
        return float(val)
    except (ValueError, TypeError):
        return None


def main():
    parser = argparse.ArgumentParser(description="Clean Air Compass ETL")
    parser.add_argument("--source", required=True, help="Path to data repo root")
    args = parser.parse_args()

    source_dir = Path(args.source)
    if not source_dir.exists():
        print(f"Error: Source directory not found: {source_dir}")
        return

    print(f"Source: {source_dir}")
    print(f"Output: {BACKEND_DATA}")
    print()

    # Clear old generated data (keep categories.json)
    for subdir in ["cities", "smogstripes", "daily", "health", "geometry"]:
        d = BACKEND_DATA / subdir
        if d.exists():
            shutil.rmtree(d)

    print("[1/7] Building master city list + composite index...")
    cities, composite, equity_data = etl_cities(source_dir)

    print("[2/7] Writing per-city detail files...")
    etl_city_details(composite, equity_data)

    print("[3/7] Processing smogstripes...")
    etl_smogstripes(source_dir)

    print("[4/7] Processing daily data...")
    etl_daily(source_dir)

    print("[5/7] Processing health impact data...")
    etl_health(source_dir)

    print("[6/7] Processing GeoJSON geometry...")
    etl_geometry(source_dir)

    print("[7/7] Copying city images...")
    etl_images(source_dir)

    print()
    print("ETL complete!")


if __name__ == "__main__":
    main()
